import pytest
import fitz
import docx
import io
from app.services.resume_parser import parse_pdf, parse_docx, parse_resume, ResumeParsingError


def create_sample_pdf_bytes() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    text = (
        "John Doe\n"
        "john.doe@example.com | +91 9876543210\n\n"
        "Education\n"
        "Master of Computer Applications (MCA), 2024\n\n"
        "Skills\n"
        "Python, FastAPI, PostgreSQL, React, Docker\n\n"
        "Experience\n"
        "Software Engineer at TechCorp\n"
    )
    page.insert_text((50, 50), text)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("Jane Smith", level=1)
    doc.add_paragraph("jane.smith@example.com | +1 555-0199")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.Tech in Computer Science")
    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph("Java, Spring Boot, MySQL, Kubernetes, Git")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Backend Developer for 3 years.")

    file_stream = io.BytesIO()
    doc.save(file_stream)
    return file_stream.getvalue()


def test_parse_pdf():
    pdf_bytes = create_sample_pdf_bytes()
    text = parse_pdf(pdf_bytes)
    assert "John Doe" in text
    assert "john.doe@example.com" in text
    assert "Python" in text


def test_parse_docx():
    docx_bytes = create_sample_docx_bytes()
    text = parse_docx(docx_bytes)
    assert "Jane Smith" in text
    assert "jane.smith@example.com" in text
    assert "Java" in text


def test_parse_resume_valid_pdf():
    pdf_bytes = create_sample_pdf_bytes()
    text = parse_resume(pdf_bytes, "resume.pdf")
    assert "John Doe" in text
    assert "PostgreSQL" in text


def test_parse_resume_valid_docx():
    docx_bytes = create_sample_docx_bytes()
    text = parse_resume(docx_bytes, "candidate_cv.docx")
    assert "Jane Smith" in text
    assert "Kubernetes" in text


def test_unsupported_file_extension():
    with pytest.raises(ResumeParsingError) as exc_info:
        parse_resume(b"some text content", "image.png")
    assert "Unsupported file format" in str(exc_info.value)


def test_doc_file_rejected():
    with pytest.raises(ResumeParsingError) as exc_info:
        parse_resume(b"legacy doc content", "legacy_resume.doc")
    assert "Unsupported file format" in str(exc_info.value)
    assert ".pdf and .docx" in str(exc_info.value)


def test_empty_file():
    with pytest.raises(ResumeParsingError) as exc_info:
        parse_resume(b"", "empty.pdf")
    assert "empty" in str(exc_info.value).lower()
